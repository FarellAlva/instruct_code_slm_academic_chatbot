"""Read and print evaluation sections of the skripsi."""
import sys, os
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

import docx

doc = docx.Document(r'd:\LLM\instruct\reference\farell_skripsi_lengkap.docx')

lines = [(i, p.style.name if p.style else "Normal", p.text.strip())
         for i, p in enumerate(doc.paragraphs) if p.text.strip()]

# Print all - focus on evaluation section
print("=== SECTION 3: EVALUASI & DATASET ===")
in_section = False
for i, style, text in lines:
    if any(kw in text.lower() for kw in [
        "dataset", "pertanyaan uji", "evaluasi", "rouge", "bleu", "bert",
        "faithfulness", "answer", "recall", "precision", "f1",
        "tabel 3", "tabel 2", "spesifikasi", "parameter inferensi",
        "kueri uji", "ground truth", "referensi jawaban"
    ]):
        print(f"\n[{i}] ({style}): {text[:500]}")

# Print tables
print("\n\n=== TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\nTable {t_idx}:")
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        if any(cells):
            print("  | " + " | ".join(cells) + " |")
    if t_idx > 10:
        break
